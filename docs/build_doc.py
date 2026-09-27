"""说明文档 Markdown -> Word 转换器（跟随官方模板结构，简洁公文排版）。"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = sys.argv[1] if len(sys.argv) > 1 else "docs/说明文档_筑安云.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else "docs/筑安云_项目技术说明文档.docx"

BLUE = RGBColor(0x1D, 0x5B, 0xD8)
GRAY = RGBColor(0x7A, 0x86, 0x9C)


def set_cjk(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_runs(p, text, size=12, name="宋体", bold=False, color=None):
    """解析 **加粗** 行内标记。"""
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_cjk(r, name, size, bold=True, color=color)
        else:
            r = p.add_run(part)
            set_cjk(r, name, size, bold=bold, color=color)


doc = Document()
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
doc.styles["Normal"].paragraph_format.line_spacing = 1.5

lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
in_code = False
code_buf = []

while i < len(lines):
    line = lines[i].rstrip()

    if line.startswith("```"):
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            r = p.add_run("\n".join(code_buf))
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            r.font.size = Pt(8.5)
            code_buf = []
        in_code = not in_code
        i += 1
        continue
    if in_code:
        code_buf.append(line)
        i += 1
        continue

    if not line.strip() or line.strip() == "---":
        i += 1
        continue

    if line.startswith("| ") or line.startswith("|-"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not set("".join(cells)) <= set("-: "):
                rows.append(cells)
            i += 1
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    c = table.cell(ri, ci)
                    c.text = ""
                    p = c.paragraphs[0]
                    add_runs(p, cell, size=10.5, bold=(ri == 0))
            doc.add_paragraph()
        continue

    if line.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        add_runs(p, line[4:], size=13, name="黑体", bold=True, color=BLUE)
    elif line.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        add_runs(p, line[3:], size=15, name="黑体", bold=True, color=BLUE)
    elif line.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, line[2:].strip() or "筑安云", size=18, name="黑体", bold=True)
    elif line.startswith("> "):
        p = doc.add_paragraph()
        add_runs(p, line[2:], size=10.5, color=GRAY)
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, line[2:], size=12)
    elif re.match(r"^\d+\. ", line):
        m = re.match(r"^(\d+)\. (.*)", line)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        add_runs(p, f"{m.group(1)}. {m.group(2)}", size=12)
    else:
        p = doc.add_paragraph()
        add_runs(p, line, size=12)

    i += 1

doc.save(OUT)
print("生成:", OUT)
