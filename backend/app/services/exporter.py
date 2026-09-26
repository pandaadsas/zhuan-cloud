"""周报导出 Word（python-docx 轻量实现：解析周报Markdown为标题/正文/列表）。"""
import io
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from ..models import WeeklyReport


def _set_cjk(style_or_run, name: str = "微软雅黑"):
    if hasattr(style_or_run, "font"):
        style_or_run.font.name = name
        rpr = style_or_run._element.get_or_add_rPr()
        rpr.rFonts.set(qn("w:eastAsia"), name)


def weekly_to_docx(report: WeeklyReport) -> bytes:
    doc = Document()
    _set_cjk(doc.styles["Normal"])
    doc.styles["Normal"].font.size = Pt(11)

    for raw in report.content_md.splitlines():
        line = raw.rstrip()
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        text = re.sub(r"\*(.+?)\*", r"\1", text)
        if not text.strip():
            continue
        if text.startswith("## "):
            p = doc.add_heading(text[3:].strip(), level=2)
            _set_cjk(p)
        elif text.startswith("# "):
            p = doc.add_heading(text[2:].strip(), level=1)
            _set_cjk(p)
        elif text.startswith("- "):
            p = doc.add_paragraph(text[2:].strip(), style="List Bullet")
            _set_cjk(p)
        elif text.startswith("---"):
            continue
        else:
            p = doc.add_paragraph(text.strip())
            _set_cjk(p)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
